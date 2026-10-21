"""Build the eight-slide ADI leadership deck from the saved comparison workbook.

Requires python-pptx and openpyxl. No API calls or credentials are used.
"""
from __future__ import annotations

import argparse
import csv
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SLIDE_COUNT = 8
WIDTH = 13.333333
HEIGHT = 7.5
FONT = "Aptos"
INK = "17212B"
MUTED = "526170"
BLUE = "2563EB"
TEAL = "007F78"
PALE_BLUE = "EFF6FF"
PALE_TEAL = "EDF8F5"
PALE_GREY = "F4F6F8"
RULE = "DCE3EA"
AMBER = "8A4B08"
PALE_AMBER = "FFF5E8"
WHITE = "FFFFFF"
METRICS = [
    ("faithfulness", "Faithfulness", "Higher"),
    ("answer_relevance", "Answer relevance", "Higher"),
    ("context_relevance", "Context relevance", "Higher"),
    ("toxicity", "Toxicity", "Lower"),
    ("tone_appropriateness", "Tone appropriateness", "Higher"),
]


def read_evidence(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables = {}
        for name in ("Summary", "Setup", "Agreement", "Latency_Cost", "Accuracy_vs_Gold", "Per_Row"):
            rows = workbook[name].iter_rows(values_only=True)
            headers = next(rows)
            tables[name] = [dict(zip(headers, row)) for row in rows if row[0] is not None]
    finally:
        workbook.close()
    evidence = {
        "summary": {row["criterion"]: row for row in tables["Summary"]},
        "setup": {row["item"]: row["value"] for row in tables["Setup"]},
        "agreement": {row["metric"]: row for row in tables["Agreement"]},
        "operations": tables["Latency_Cost"],
        "accuracy": {(row["metric"], row["judge"]): row for row in tables["Accuracy_vs_Gold"]},
        "examples": {int(row["row_id"]): row for row in tables["Per_Row"] if row["row_id"] in (2, 8)},
    }
    if evidence["setup"]["Runs per judge"] != 2:
        raise ValueError("This deck describes the saved two-run benchmark.")
    if evidence["setup"]["Concurrency (both judges)"] != 3:
        raise ValueError("This deck describes concurrency 3; update the narrative for another setup.")
    if any(row["rows"] != 35 for row in evidence["operations"]):
        raise ValueError("This deck describes the 35-row exploratory benchmark.")
    if len(evidence["operations"]) != 4:
        raise ValueError("Expected two runs for each of the two measured judges.")
    if evidence["setup"]["Pass/fail threshold on 0-1 scores"] != 0.5:
        raise ValueError("The worked examples describe the benchmark's 0.5 cut-off.")
    for key, _, _ in METRICS:
        if evidence["agreement"][key]["n"] != 35:
            raise ValueError(f"Missing comparable results for {key}.")
    with (HERE / "data" / "synthetic_adi_gt_35.csv").open(encoding="utf-8-sig", newline="") as source:
        cases = list(csv.DictReader(source))
    evidence["cases"] = {row_id: cases[row_id] for row_id in (2, 8)}
    for row_id, case in evidence["cases"].items():
        if case["User Query"] != evidence["examples"][row_id]["user_query"]:
            raise ValueError(f"Source query and report example do not align for row_id {row_id}.")
    atm = evidence["examples"][2]
    vpn = evidence["examples"][8]
    if not (atm["gold_faithfulness"] == 0 and atm["jev_faithfulness"] < 0.5 <= atm["llm_faithfulness"]):
        raise ValueError("The ATM case no longer supports the stated faithfulness contrast.")
    if not (vpn["gold_answer_correctness"] == vpn["llm_answer_correctness"] == 1
            and round(vpn["jev_answer_correctness"] * 2) / 2 == 0.5
            and vpn["llm_tone_appropriateness"] == vpn["jev_tone_appropriateness"] == 0):
        raise ValueError("The VPN case no longer supports the stated correctness/tone contrast.")
    if ("Rs. 3,00,000" not in evidence["cases"][2]["Content Retrieved by ADI"]
            or "Rs. 1,00,000" not in evidence["cases"][2]["Response by ADI"]):
        raise ValueError("The source ATM limits have changed; revise the condensed example.")
    return evidence


def colour(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rectangle(slide, x, y, width, height, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour(fill)
    shape.line.fill.background()
    return shape


def text(slide, x, y, width, height, content, size=18, bold=False,
         fill=INK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = colour(fill)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(4)
    return shape


def table(slide, x, y, widths, headers, rows, row_heights, size=16,
          numeric_columns=()):
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y),
        Inches(sum(widths)), Inches(sum(row_heights))
    )
    grid = shape.table
    for index, width in enumerate(widths):
        grid.columns[index].width = Inches(width)
    for row_index, values in enumerate([headers] + rows):
        grid.rows[row_index].height = Inches(row_heights[row_index])
        for column_index, value in enumerate(values):
            cell = grid.cell(row_index, column_index)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Inches(0.13)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = colour(
                BLUE if row_index == 0 else
                PALE_TEAL if column_index == len(headers) - 1 else
                WHITE if row_index % 2 else PALE_GREY
            )
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = FONT
                paragraph.font.size = Pt(size)
                paragraph.font.bold = row_index == 0 or column_index == 0
                paragraph.font.color.rgb = colour(WHITE if row_index == 0 else INK)
                paragraph.space_before = paragraph.space_after = Pt(0)
                if column_index in numeric_columns:
                    paragraph.alignment = PP_ALIGN.CENTER
    return shape


def base_slide(deck, number, title, source, benchmark_date):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colour(WHITE)
    rectangle(slide, 0, 0, WIDTH, 0.06, BLUE)
    text(slide, 0.58, 0.26, 4.5, 0.24, "ADI / MODEL EVALUATION", 11, True, BLUE)
    text(slide, 7.1, 0.26, 5.65, 0.24,
         f"EXPLORATORY ASSESSMENT  |  {benchmark_date}", 10, fill=MUTED,
         align=PP_ALIGN.RIGHT)
    text(slide, 0.58, 0.78, 12.17, 0.61, title, 29, True)
    rectangle(slide, 0.58, 7.02, 12.17, 0.012, RULE)
    text(slide, 0.58, 7.13, 11.25, 0.2, source, 9.5, fill=MUTED)
    text(slide, 12.0, 7.10, 0.75, 0.25, f"{number:02d} / {SLIDE_COUNT:02d}", 11, True,
         fill=BLUE, align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


def summary_value(evidence, criterion, judge):
    return float(evidence["summary"][criterion][judge])


def build_deck(evidence: dict, output: Path, benchmark_date: date) -> None:
    deck = Presentation()
    deck.slide_width = Inches(WIDTH)
    deck.slide_height = Inches(HEIGHT)
    deck.core_properties.title = "ADI: JEV as an exploratory evaluation alternative"
    deck.core_properties.subject = "Leadership review of the saved JEV vs GPT-4o judge benchmark"
    deck.core_properties.author = "ADI"
    deck.core_properties.keywords = "ADI, JEV, GPT-4o, evaluation, exploratory benchmark"
    shown_date = benchmark_date.strftime("%d %b %Y").upper()
    cost_llm = summary_value(evidence, "Cost per 1,000 rows (USD)", "LLM judge")
    cost_jev = summary_value(evidence, "Cost per 1,000 rows (USD)", "JEV")
    throughput_llm = summary_value(evidence, "Throughput (rows/s)", "LLM judge")
    throughput_jev = summary_value(evidence, "Throughput (rows/s)", "JEV")
    accuracy_llm = summary_value(evidence, "Accuracy vs gold (mean of 5 metrics)", "LLM judge")
    accuracy_jev = summary_value(evidence, "Accuracy vs gold (mean of 5 metrics)", "JEV")

    slide = base_slide(deck, 1, "JEV: a lower-cost evaluation option for ADI",
                       "Source: comparison workbook / Summary and Setup. Synthetic evidence, not a production assessment.", shown_date)
    text(slide, 0.58, 1.55, 12.17, 0.85,
         "Pilot JEV for bounded scoring and classification;\nkeep GPT-4o for complex review and explanations.", 22, True)
    rectangle(slide, 0.58, 2.55, 12.17, 1.26, PALE_BLUE)
    kpis = [
        (f"{cost_llm / cost_jev:.1f}x", "Lower estimated token cost", "JEV vs GPT-4o judge"),
        (f"{throughput_jev / throughput_llm:.1f}x", "Higher dataset throughput", "Same concurrency; batching differs"),
        (f"{accuracy_jev:.1%}", "JEV quality-label accuracy", f"GPT-4o: {accuracy_llm:.1%}; designer labels"),
    ]
    for index, (value, label, detail) in enumerate(kpis):
        xpos = 0.82 + index * 4.04
        text(slide, xpos, 2.67, 3.6, 0.56, value, 34, True, BLUE)
        text(slide, xpos, 3.25, 3.75, 0.29, label, 16, True)
        text(slide, xpos, 3.55, 3.75, 0.19, detail, 11.5, fill=MUTED)
    text(slide, 0.58, 4.16, 5.8, 0.3, "WHAT WAS TESTED", 14, True, TEAL)
    text(slide, 0.58, 4.58, 5.8, 1.45,
         "35 synthetic ADI cases; 2 runs per judge.\nFive quality metrics + a correctness check.\nConcurrency 3; context cap 8,000 characters.", 18)
    text(slide, 6.88, 4.16, 5.87, 0.3, "WHAT IS NOT ESTABLISHED", 14, True, AMBER)
    text(slide, 6.88, 4.58, 5.87, 1.45,
         "GPT-4o-mini and SLMs were not tested.\nNo proof of replacing answer generation.\nNo production uptime or compliance validation.", 18)
    rectangle(slide, 0.58, 6.37, 12.17, 0.43, PALE_AMBER)
    text(slide, 0.76, 6.45, 11.81, 0.25,
         "Decision requested: approve an independently labelled validation pilot, not a model replacement.", 15, True, AMBER)
    notes(slide,
          "Executive positioning: JEV is a promising alternative for bounded evaluation and structured decisions. "
          "This experiment did not test ADI answer generation, GPT-4o-mini or any SLM. "
          "Measured models: openai/gpt-4o via OpenRouter and JEV jev-1.13.0 via TypeSafe. "
          "The 35 synthetic rows were run twice per judge. Quality-label accuracy is mean binary accuracy "
          "over five metrics against designer-created labels: GPT-4o 94.6%, JEV 98.6%. "
          "Labels and JEV rubrics were designed by the same author; this can bias results. "
          "Cost and throughput ratios are calculated from the saved workbook, not vendor marketing claims. "
          "Token cost excludes infrastructure, implementation, human review and any escalation model. "
          "Use the findings to justify a pilot, not to claim general model superiority.")

    slide = base_slide(deck, 2, "Different model roles, not just different prices",
                       "Sources: implemented evaluator APIs and workbook Setup. GPT-4o-mini / SLM column describes options, not results.", shown_date)
    text(slide, 0.58, 1.48, 12.17, 0.38,
         "JEV offers structured decisions; generative models retain broader language output and review capability.", 17, fill=MUTED)
    table(slide, 0.58, 2.03, [2.0, 3.27, 3.35, 3.55],
          ["Dimension", "GPT-4o judge", "GPT-4o-mini / SLM", "JEV"], [
              ["Core approach", "Generates JSON scores\nand written reasons", "Generative judging; SLM\ndepends on chosen model", "Typed Score decisions;\nprobabilities + confidence"],
              ["Integration", "Prompt + query, answer\nand retrieved context", "Prompt / tuning;\nvalidation required", "State + rubric questions;\n5 quality metrics per call"],
              ["Best-fit role", "Complex review;\nexplanations; escalation", "Candidate low-cost judge\nor answer-generation model", "Routine scoring; guardrail /\nrouting candidates to test"],
              ["Trade-off", "Flexible; explanation cost;\nprompt / parsing sensitivity", "Quality and cost depend on\nmodel, hosting and tuning", "No generated rationale;\nrubric / threshold sensitivity"],
              ["Evidence today", "35 cases x 2 runs", "Not benchmarked;\nno ranking supported", "35 cases x 2 runs;\nmodel jev-1.13.0"],
          ], [0.5] + [0.76] * 5, size=16)
    text(slide, 0.58, 6.53, 12.17, 0.3,
         "SLM = small language model, not one specific model. Select the model and deployment before comparing it.", 13, fill=MUTED)
    notes(slide,
          "The comparison is between judging workflows, not between end-to-end chatbot answer models. "
          "GPT-4o generates strict JSON with metric scores and brief written reasons. JEV Score returns "
          "a distribution over rubric levels, an expected raw score and confidence, without generated reasons. "
          "The tested JEV integration submits one row with five quality questions and an optional correctness question. "
          "Routing and guardrail uses are proposed adjacent experiments, not measured benefits in this benchmark. "
          "GPT-4o-mini and an SLM may also be candidates, but no accuracy, price, latency or reliability comparison "
          "against them has been run here. An SLM deployment must include hosting, tuning and operating costs. "
          "Reference: https://docs.typesafe.ai/api.md and https://docs.typesafe.ai/primitives/score.md")

    slide = base_slide(deck, 3, "Same 0-1 range, different scoring mathematics",
                       "Sources: adi_gt_evaluator.py metric prompts; jev_judge.py rubrics and score normalization.", shown_date)
    text(slide, 0.58, 1.48, 12.17, 0.3,
         "Quality metrics use the query, answer and context. Correctness additionally needs the SME's corrected answer.", 16, fill=MUTED)
    table(slide, 0.58, 1.98, [2.3, 5.12, 4.75],
          ["Metric", "GPT-4o: prompt-directed scoring", "JEV: probabilistic rubric scoring"], [
              ["Faithfulness", "Supported atomic claims / all factual claims.\nNo factual claims -> 1.0.", "4 support levels (0-3);\nexpected level / 3."],
              ["Answer relevance", "On-topic response statements /\nall response statements; ignore accuracy.", "4 on-topic levels (0-3);\nexpected level / 3."],
              ["Context relevance", "Relevant retrieved-context statements /\nall retrieved-context statements.", "4 retrieval-relevance levels (0-3);\nexpected level / 3."],
              ["Toxicity", "G-Eval raw 1-5 -> (raw - 1) / 4.\nLower = safer.", "4 severity levels (0-3);\nexpected level / 3; lower = safer."],
              ["Tone", "G-Eval raw 1-5 -> (raw - 1) / 4.\nHigher = more professional.", "4 professional-tone levels (0-3);\nexpected level / 3."],
              ["Correctness\n(GT required)", "Semantic match to SME correction;\n0 / 0.5 / 1 on applicable rows.", "3 semantic-match levels (0-2);\nexpected level / 2 on applicable rows."],
          ], [0.46] + [0.64] * 6, size=16)
    text(slide, 0.58, 6.41, 12.17, 0.26,
         "JEV raw = sum(level x probability). Example: [0.10, 0.20, 0.30, 0.40] -> 2.00 / 3 = 0.667.", 14, True, TEAL)
    text(slide, 0.58, 6.73, 12.17, 0.21,
         "GPT ratios are model estimates, not independently counted claims. JEV confidence is separate from the score.", 12.5, fill=MUTED)
    notes(slide,
          "Important: these scores are not mathematically equivalent, despite sharing the 0-1 range. "
          "The GPT prompt asks for atomic claim or statement ratios, but the application does not independently "
          "extract and count claims. It accepts the model's returned judgment. Toxicity and tone use G-Eval "
          "five-point ratings normalized by (raw-1)/4. These are RAGAS-style prompts, not a run of the RAGAS library. "
          "JEV returns an expected ordinal rubric score. For four criteria indexed 0 to 3, "
          "normalized score = sum(k * probability[k]) / 3. Example: 0*.10 + 1*.20 + 2*.30 + 3*.40 = 2.00, "
          "then 2/3 = .667. The implementation normalizes the returned answer['score']; it does not use "
          "confidence as the quality score. The raw expected score can be fractional. "
          "Correctness uses three levels and divides by 2. It is separately evaluated on the 15 SME-No rows "
          "with a corrected answer. The standalone evaluators automatically give SME-Yes rows correctness 1.0; "
          "the comparison excludes those automatic credits. Product/subproduct matches and document "
          "precision/recall are deterministic ground-truth checks, not alternate model-scoring methods. "
          "Rubric and threshold calibration must be validated per metric. These scores are not established "
          "calibrated probabilities of error.")

    slide = base_slide(deck, 4, "Five quality metrics without ground-truth answers",
                       "Source: comparison workbook / Agreement. Means and verdict agreement averaged over 2 runs, 35 cases each.", shown_date)
    text(slide, 0.58, 1.48, 12.17, 0.3,
         "These metrics need no SME reference answer. Inter-judge agreement is not a measure of accuracy.", 17, fill=MUTED)
    metric_rows = []
    for key, label, direction in METRICS:
        result = evidence["agreement"][key]
        metric_rows.append([
            label, direction, f"{result['mean_llm']:.4f}",
            f"{result['mean_jev']:.4f}", f"{result['verdict_agreement']:.1%}",
        ])
    table(slide, 0.58, 2.06, [2.7, 1.37, 2.25, 2.25, 3.6],
          ["Quality metric", "Better", "GPT-4o mean", "JEV mean", "Pass/fail agreement"],
          metric_rows, [0.55] + [0.61] * 5, size=17, numeric_columns=(2, 3, 4))
    agreement = evidence["agreement"]["faithfulness"]["verdict_agreement"]
    rectangle(slide, 0.58, 6.00, 12.17, 0.48, PALE_AMBER)
    text(slide, 0.76, 6.10, 11.81, 0.27,
         f"Faithfulness diverges most ({agreement:.1%} agreement): review unsupported-claim disagreements.", 16, True, AMBER)
    text(slide, 0.58, 6.64, 12.17, 0.29,
         "Cut-off = 0.5; high toxicity fails. Excludes correctness, SME agreement, tag matches and document precision/recall.", 12.5, fill=MUTED)
    notes(slide,
          "This is the requested separate non-ground-truth metric comparison. It deliberately excludes "
          "answer correctness, SME agreement, product/subproduct matches and document retrieval precision/recall. "
          "Faithfulness still needs the retrieved context as evidence, but no independently corrected SME answer. "
          "Scores are means from the workbook Agreement sheet, averaged over both benchmark runs. "
          "Pass/fail agreement uses the shared .5 threshold; high toxicity is a failure. "
          "Mean score differences are descriptive and must not be interpreted as one model generating better "
          "answers; both judges scored identical ADI outputs. Different scoring mathematics can explain "
          "different numeric averages. Faithfulness had only 77.1% verdict agreement, so unsupported-claim "
          "cases are a priority for SME adjudication. Accuracy claims elsewhere in this deck do require "
          "designer labels; this table's metrics themselves do not require ground-truth reference answers.")

    slide = base_slide(deck, 5, "Where each judge trails: the measured trade-offs",
                       "Source: workbook / Accuracy_vs_Gold. Designer labels; means of 2 runs. Sample size and measure vary by row.", shown_date)
    text(slide, 0.58, 1.48, 12.17, 0.35,
         "JEV trails on correctness grades; GPT-4o trails on unsupported-claim detection and toxicity classification.", 17, fill=MUTED)
    gap_rows = []
    comparisons = [
        ("answer_correctness", "Answer correctness\nBand accuracy (n=15)", "accuracy"),
        ("faithfulness", "Faithfulness\nFailure recall (10 unsupported cases)", "recall"),
        ("toxicity", "Toxicity\nVerdict accuracy (n=35)", "accuracy"),
        ("answer_relevance", "Answer relevance\nVerdict accuracy (n=35)", "accuracy"),
        ("context_relevance", "Context relevance\nVerdict accuracy (n=35)", "accuracy"),
        ("tone_appropriateness", "Tone appropriateness\nVerdict accuracy (n=35)", "accuracy"),
    ]
    for metric, label, measure in comparisons:
        gpt_value = float(evidence["accuracy"][(metric, "LLM judge")][measure])
        jev_value = float(evidence["accuracy"][(metric, "JEV")][measure])
        difference = abs(gpt_value - jev_value) * 100
        leader = "GPT-4o" if gpt_value > jev_value else "JEV"
        comparison = "Tie on verdict accuracy" if difference < 0.01 else f"{leader} ahead by {difference:.1f} pp"
        gap_rows.append([label, f"{gpt_value:.1%}", f"{jev_value:.1%}", comparison])
    table(slide, 0.58, 2.06, [4.25, 1.65, 1.65, 4.62],
          ["Metric / success measure", "GPT-4o", "JEV", "Observed gap"], gap_rows,
          [0.48] + [0.65] * 6, size=16, numeric_columns=(1, 2))
    text(slide, 0.58, 6.59, 12.17, 0.32,
         "pp = percentage points. Small synthetic sample; rubric and threshold differences limit generalization.", 13.5, True, AMBER)
    notes(slide,
          "This slide compares judge success against the synthetic designer labels, not average quality-score levels. "
          "Unlike slide 4, this is a labelled validation view. It includes the ground-truth-dependent correctness "
          "metric separately and names the measure used for each row. "
          "JEV lags GPT-4o on answer-correctness band accuracy: 73.3% versus 80.0%, a 6.7-percentage-point gap "
          "on only 15 applicable cases. That is one additional case per run and is not statistically established superiority. "
          "Correctness scores are rounded to the nearest 0, 0.5 or 1 before comparison with the designer grade. "
          "GPT-4o lags JEV on detecting unsupported factual claims at the chosen 0.5 cut-off: "
          "50% versus 100% recall on ten designer-unfaithful cases. Faithfulness verdict accuracy is also "
          "78.6% versus 95.7%. In a case-specific review, GPT may notice a wrong fact in its written reason "
          "yet still return a score that passes the downstream cut-off. "
          "Toxicity verdict accuracy is 97.1% for GPT versus 100% for JEV, representing one additional false "
          "positive from GPT on slang. There are only two designer-toxic cases, so this is a limited signal. "
          "Answer relevance and tone accuracy tie at 100%; context relevance ties at 97.1%. "
          "Do not interpret equal verdict accuracy as identical raw scores or equally calibrated probabilities. "
          "All results are means of the two saved runs; no model calls were rerun to select these examples. "
          "The labels and rubrics were designed by the same author. These are workflow-level observations "
          "under the implemented rubrics and thresholds, not broad rankings of model intelligence.")

    slide = base_slide(deck, 6, "Two worked examples: one advantage for each judge",
                       "Sources: workbook / Per_Row (run 1); synthetic cases 03 and 09, row_id 2 and 8. Full text in presenter notes.", shown_date)
    text(slide, 0.58, 1.48, 12.17, 0.34,
         "Both judges evaluated the same ADI answers. The advantage is metric-specific, not an overall answer approval.", 16.5, fill=MUTED)
    atm = evidence["examples"][2]
    vpn = evidence["examples"][8]
    example_panels = [
        {
            "row_id": 2, "x": 0.58, "fill": PALE_TEAL,
            "heading": "CASE 03: JEV BETTER ON FAITHFULNESS",
            "reference": "ATM limit: Rs. 3,00,000.\nPOS / e-commerce limit: Rs. 6,00,000.",
            "answer": "ATM limit: Rs. 1,00,000.\nPOS / e-commerce limit: Rs. 6,00,000.",
            "scores": [f"{atm['llm_faithfulness']:.2f} / not flagged", f"{atm['jev_faithfulness']:.2f} / flagged"],
            "expected": "Reference verdict: unsupported claim; should be flagged.",
            "why": "JEV flags the wrong limit. GPT's reason spots it, but its score still passes the 0.5 cut-off.",
        },
        {
            "row_id": 8, "x": 6.865, "fill": PALE_BLUE,
            "heading": "CASE 09: GPT-4o BETTER ON CORRECTNESS",
            "reference": "Reset via SSPR + registered-mobile OTP;\nnew password syncs to VPN within 15 minutes.",
            "answer": "SSPR + mobile OTP; sync in 15 minutes.\nThe response also uses rude phrasing.",
            "scores": [f"{vpn['llm_answer_correctness']:.2f} / full match", f"{vpn['jev_answer_correctness']:.2f} / partial"],
            "expected": "Reference: full factual match; both judges reject the tone.",
            "why": "GPT matches the reference grade. JEV falls just below the 0.75 full-match band; confidence is low.",
        },
    ]
    for panel in example_panels:
        xpos = panel["x"]
        rectangle(slide, xpos, 2.00, 5.885, 0.40, panel["fill"])
        text(slide, xpos + 0.14, 2.10, 5.605, 0.24, panel["heading"], 14, True)
        text(slide, xpos + 0.14, 2.56, 5.605, 0.66,
             evidence["cases"][panel["row_id"]]["User Query"], 17, True)
        text(slide, xpos + 0.14, 3.29, 5.605, 0.20,
             "REFERENCE (CONDENSED)", 11, True, TEAL)
        text(slide, xpos + 0.14, 3.55, 5.605, 0.59, panel["reference"], 16)
        text(slide, xpos + 0.14, 4.21, 5.605, 0.20,
             "ADI ANSWER (CONDENSED)", 11, True, BLUE)
        text(slide, xpos + 0.14, 4.47, 5.605, 0.63, panel["answer"], 16)
        table(slide, xpos, 5.19, [2.9425, 2.9425], ["GPT-4o score", "JEV score"],
              [panel["scores"]], [0.35, 0.47], size=15.5, numeric_columns=(0, 1))
        text(slide, xpos + 0.14, 6.13, 5.605, 0.27, panel["expected"], 13, True)
        text(slide, xpos + 0.14, 6.48, 5.605, 0.42, panel["why"], 12.5, fill=MUTED)
    example_notes = [
        "The two selected examples illustrate metric-specific differences under the current scoring pipeline, "
        "not that one judge is right about every aspect of either answer. Both differences persist in both runs. "
        "Visible scores are run 1. Case numbering is one-based; source row_ids are zero-based. "
        "Case 03 / row_id 2: JEV faithfulness .33 is below .5 and is flagged; GPT .67 passes. "
        "GPT's written reason is 'ATM limit claim not supported', so the error is score/threshold behaviour, "
        "not proof that GPT could not identify the contradiction. In run 2, GPT faithfulness is .50 "
        "and JEV is .98/3 = .3267, so the same verdict contrast remains. "
        "Case 09 / row_id 8: the designer's factual correctness grade is 1.0, but the SME marked the overall "
        "response No because it uses rude, belittling phrasing. GPT returns correctness 1.0 with reason "
        "'Semantically equivalent.' JEV returns raw 1.44/2 = .72, which is mapped to the partial band .5. "
        "Both runs give the same correctness scores. Both judges give tone 0, so neither should approve "
        "the overall answer. JEV correctness confidence is .34; it is near the .75 full-match boundary. "
        "Its probabilities are .56 on partial and .44 on full match, so it is uncertain rather than "
        "confidently asserting a factual contradiction. JEV provides no written reason; do not attribute "
        "its lower score to a specific internal cause. The slide's reference and answer blocks are condensed "
        "for readability; the original query, retrieved context, ADI answer and corrected answer follow."
    ]
    for row_id in (2, 8):
        case = evidence["cases"][row_id]
        example_notes.append(f"\nCASE {row_id + 1:02d} / row_id {row_id}")
        for column in ("User Query", "Content Retrieved by ADI", "Response by ADI",
                       "If ADI response incorrect(Enter right response)"):
            example_notes.append(f"{column}: {case[column]}")
    notes(slide, "\n".join(example_notes))

    slide = base_slide(deck, 7, "Measured economics and reliability",
                       "Sources: comparison workbook / Summary, Latency_Cost and Consistency. USD token estimates; same local test setup.", shown_date)
    text(slide, 0.58, 1.48, 12.17, 0.32,
         "35 rows, 2 runs, concurrency 3. Token estimates are not total cost; observed stability is not a service SLA.", 16.5, fill=MUTED)
    criteria = [
        ("Cost per 1,000 evaluated rows", "Cost per 1,000 rows (USD)", "${:.3f}"),
        ("Wall time for 35-row dataset", "Wall time for the dataset (s)", "{:.2f} s"),
        ("Dataset throughput", "Throughput (rows/s)", "{:.2f} rows/s"),
        ("Median API-call latency", "Median call latency (s)", "{:.3f} s"),
        ("p95 API-call latency", "p95 call latency (s)", "{:.3f} s"),
    ]
    operation_rows = [
        [label, pattern.format(summary_value(evidence, key, "LLM judge")),
         pattern.format(summary_value(evidence, key, "JEV"))]
        for label, key, pattern in criteria
    ]
    totals = {}
    for judge in ("LLM judge", "JEV"):
        runs = [row for row in evidence["operations"] if row["judge"] == judge]
        totals[judge] = {
            key: sum(int(row[key]) for row in runs)
            for key in ("api_calls", "failed_calls", "rows", "rows_missing_scores")
        }
    operation_rows.extend([
        ["Failed calls / total observed (both runs)"] + [
            f"{totals[judge]['failed_calls']} / {totals[judge]['api_calls']}"
            for judge in ("LLM judge", "JEV")],
        ["Rows with missing scores / total (both runs)"] + [
            f"{totals[judge]['rows_missing_scores']} / {totals[judge]['rows']}"
            for judge in ("LLM judge", "JEV")],
        ["Mean per-metric verdict flip rate"] + [
            f"{summary_value(evidence, 'Run-to-run verdict flip rate', judge):.2%}"
            for judge in ("LLM judge", "JEV")],
        ["Mean absolute run-to-run score change"] + [
            f"{summary_value(evidence, 'Run-to-run mean score change', judge):.4f}"
            for judge in ("LLM judge", "JEV")],
    ])
    table(slide, 0.58, 2.03, [6.35, 2.91, 2.91],
          ["Measure", "GPT-4o judge", "JEV"], operation_rows,
          [0.44] + [0.41] * 9, size=16, numeric_columns=(1, 2))
    text(slide, 0.58, 6.32, 12.17, 0.25,
         "Batching differs: GPT-4o uses 5 rows/call + correctness calls; JEV uses 1 row/call. Prefer dataset throughput.", 13, True, AMBER)
    text(slide, 0.58, 6.65, 12.17, 0.25,
         "Rates per 1M tokens: GPT-4o $2.50 input / $10 output; JEV $0.042 input / $0 output. No uptime claim.", 12.5, fill=MUTED)
    notes(slide,
          "All latency, cost and throughput values come from the saved live benchmark. Aggregate rows show "
          "the mean of the two run-level statistics; the p95 is the mean of each run's p95, not a pooled p95. "
          "GPT-4o made 10 calls per run: 7 batched quality calls plus 3 separate correctness calls. "
          "JEV made 35 calls per run, one row with parallel Score questions in each request. "
          "This makes raw call latency an unequal-workload comparison. The 35-row wall time and dataset "
          "throughput are more decision-relevant, although batching and providers still differ. "
          "Zero failed calls and zero missing-score rows were observed across 20 GPT calls and 70 JEV calls. "
          "No rate-limit, outage, recovery or long-duration service test was performed. "
          "Run-to-run metrics are means across the six scored metrics, including applicable correctness rows. "
          "They measure repeatability, not correctness or calibrated confidence. "
          "Cost uses successful usage tokens at the recorded token rates; it excludes infrastructure, "
          "engineering, review and any future GPT/SME escalation. The extrapolation per 1,000 rows assumes "
          "the same input lengths, output lengths, batching and correctness mix. "
          "A future GPT-4o-mini/SLM comparison must measure total cost and the same quality gates.")

    slide = base_slide(deck, 8, "ADI recommendation: a confidence-gated pilot",
                       "Recommendation, not a measured production architecture. Correctness figures: workbook Summary, 15 applicable cases.", shown_date)
    text(slide, 0.58, 1.50, 12.17, 0.42, "Approve validation, not replacement.", 23, True, BLUE)
    stages = [
        (0.58, 3.22, PALE_TEAL, "JEV routine screening"),
        (4.12, 4.72, PALE_AMBER, "Low confidence / disagreement / high risk"),
        (9.16, 3.59, PALE_BLUE, "GPT-4o or SME review"),
    ]
    for xpos, width, fill, label in stages:
        rectangle(slide, xpos, 2.20, width, 0.86, fill)
        label_shape = text(slide, xpos + 0.13, 2.33, width - 0.26, 0.60, label, 17, True,
                           align=PP_ALIGN.CENTER)
        label_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for xpos in (3.87, 8.91):
        chevron = slide.shapes.add_shape(
            MSO_SHAPE.CHEVRON, Inches(xpos), Inches(2.51), Inches(0.16), Inches(0.24)
        )
        chevron.fill.solid()
        chevron.fill.fore_color.rgb = colour(TEAL)
        chevron.line.fill.background()
    table(slide, 0.58, 3.24, [2.07, 5.18, 4.92],
          ["Phase", "Recommended action", "Decision evidence"], [
              ["1. Shadow test", "500+ de-identified, independently SME-labelled\ncases; include numbers, refusals and Hinglish.", "Failure recall, correctness, latency;\nconfidence slices and vendor data controls."],
              ["2. Bounded pilot", "Start with routine quality scoring;\nkeep high-risk answers in the escalation path.", "Calibrate score / confidence cut-offs;\nretain audit logs and monitor drift."],
              ["3. Compare options", "Test GPT-4o-mini + named SLMs on the same\ndata, task definitions and operating budgets.", "Per-metric recall, explanations, privacy,\nhosting, availability and total cost."],
          ], [0.46] + [0.79] * 3, size=16)
    correctness_llm = summary_value(evidence, "Answer-correctness band accuracy", "LLM judge")
    correctness_jev = summary_value(evidence, "Answer-correctness band accuracy", "JEV")
    text(slide, 0.58, 6.31, 12.17, 0.63,
         f"Retain GPT-4o for difficult correctness review: {correctness_llm:.1%} vs JEV {correctness_jev:.1%} band accuracy\n"
         "on 15 applicable synthetic cases. Human review remains necessary for high-risk banking decisions.", 15.5, True, AMBER)
    notes(slide,
          "Proposed leadership decision: authorize a controlled shadow-mode pilot with data-security approval. "
          "500+ de-identified cases is a proposed pilot scope, not a claim that the current benchmark contains "
          "500 cases or that 500 suffices for all rare risks. Use independently SME-labelled held-out cases, "
          "including difficult numerical, eligibility, multilingual and refusal examples. Separate data used "
          "to tune rubrics or thresholds from final validation. Report confidence intervals and per-metric "
          "false negatives rather than relying on one overall accuracy number. "
          "Keep high-risk cases in the GPT/SME path irrespective of confidence. JEV's confidence is not "
          "yet calibrated for ADI; the small high-confidence subsets had no errors but provide no guarantee. "
          "Validate any confidence threshold on held-out ADI cases before gating decisions. "
          "Review vendor retention, residency, access controls, contractual terms and operational availability "
          "before transmitting banking data. Maintain existing deterministic SME/document checks. "
          "GPT-4o achieved 80.0% vs JEV 73.3% band accuracy on 15 applicable correctness cases; this small "
          "difference is directional, not statistically established superiority. "
          "Keep GPT-4o for written rationale and complex adjudication during the pilot. "
          "Add GPT-4o-mini and explicitly selected hosted/self-hosted SLMs to the same held-out benchmark. "
          "Production expansion requires agreed quality, compliance and total-cost gates; answer-generation "
          "replacement requires an entirely separate experiment.")

    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    validate_deck(output)


def validate_deck(path: Path) -> None:
    deck = Presentation(path)
    if len(deck.slides) != SLIDE_COUNT:
        raise ValueError(f"Expected exactly {SLIDE_COUNT} slides.")
    for index, slide in enumerate(deck.slides, start=1):
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError(f"Slide {index} is missing presenter notes.")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                raise ValueError(f"Slide {index}: negative shape position.")
            if shape.left + shape.width > deck.slide_width + 10:
                raise ValueError(f"Slide {index}: shape extends beyond right edge.")
            if shape.top + shape.height > deck.slide_height + 10:
                raise ValueError(f"Slide {index}: shape extends beyond bottom edge.")
    print(f"Validated: {len(deck.slides)} slides, presenter notes, on-slide shape bounds.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path,
                        default=HERE / "results" / "jev_vs_llm_comparison.xlsx")
    parser.add_argument("--out", type=Path,
                        default=HERE / "results" / "ADI_JEV_Leadership_Review_v2.pptx")
    parser.add_argument("--date", type=date.fromisoformat, default=date(2026, 10, 12),
                        help="Benchmark date, YYYY-MM-DD (not the generation date).")
    args = parser.parse_args()
    build_deck(read_evidence(args.report), args.out, args.date)
    print(f"Created: {args.out}")


if __name__ == "__main__":
    main()