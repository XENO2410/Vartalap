"""Insert the three requested sections into the leadership Word report.

Sections added (before Appendix A):
  11. How the two judges work in this project (flow diagram)
  12. Worked example, step by step: JEV vs GPT on one unfaithful answer
  13. How GPT and JEV work under the hood (beginner friendly)
"""
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX = r"c:\Users\tarun\Downloads\Doc.docx"
BACKUP = r"c:\Users\tarun\Downloads\Doc_backup_20261027.docx"
PNG = r"d:\Github\ADI\jev_judge_eval\results\flow_jev_vs_gpt.png"

shutil.copy2(DOCX, BACKUP)
print("backup ->", BACKUP)

doc = Document(DOCX)

# Locate the Appendix A heading (insertion anchor).
anchor = None
for p in doc.paragraphs:
    if p.style.name.startswith("Heading") and p.text.strip().startswith("Appendix A"):
        anchor = p._p
        break
if anchor is None:
    raise SystemExit("Appendix A heading not found")
body = doc.element.body
anchor_idx = body.index(anchor)

created = []


def add_heading(text, level=2, page_break=False):
    p = doc.add_heading(text, level=level)
    if page_break:
        p.paragraph_format.page_break_before = True
    created.append(p._p)
    return p


def add_para(runs, italic=False, size=None, color=None, center=False):
    """runs: list of (text, bold) tuples."""
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, bold in runs:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    created.append(p._p)
    return p


def add_bullet(runs):
    p = doc.add_paragraph(style="List Bullet")
    for text, bold in runs:
        r = p.add_run(text)
        r.bold = bold
    created.append(p._p)
    return p


def add_table(header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    for i, htext in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(htext)
        r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    created.append(t._tbl)
    return t


def add_picture(path, width_in=6.3):
    doc.add_picture(path, width=Inches(width_in))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    created.append(p._p)
    return p


# ============================================================
# 11. Flow diagram
# ============================================================
add_heading("11. How the two judges work in this project (flow diagram)", page_break=True)
add_para([
    ("Both judges receive exactly the same row and must produce the same 0-1 metrics. "
     "The difference is how each one gets to the number: GPT-4o ", False),
    ("writes", True),
    (" a short structured answer and our code reads the numbers out of it, while JEV ", False),
    ("ticks fixed answer levels", True),
    (" and returns probabilities - there is no free text to parse.", False),
])
add_picture(PNG)
add_para([("Figure: the end-to-end judging pipeline used in this study. Same state in, "
           "same 0-1 metrics out.", False)], italic=True, size=9,
         color=RGBColor(0x59, 0x59, 0x59), center=True)
add_para([("Step by step:", True)])
add_para([
    ("1. One input row. ", True),
    ("A row is the user query, the ADI response, the retrieved context (full text plus "
     "ordered chunks), any prior conversation, and the SME's Yes/No verdict with the "
     "corrected answer where applicable.", False),
])
add_para([
    ("2. Same preparation for both judges. ", True),
    ("prepare_rows builds the identical state for every arm, so no judge gets extra "
     "information. Fairness is structural, not promised.", False),
])
add_para([
    ("3. GPT path. ", True),
    ("The prompt asks GPT to list the factual claims and statements and judge each one "
     "against the context, the query and the reference answer. GPT writes a JSON answer. "
     "Our code then does all the arithmetic (claim ratios, factual F1, the 75/25 "
     "correctness blend) - the model never calculates anything itself.", False),
])
add_para([
    ("4. JEV path. ", True),
    ("One HTTP request carries the row's state plus one typed question per metric, each "
     "with fixed ordered answer levels. JEV returns a probability for every level plus a "
     "confidence. Code converts the expected level to 0-1 and reuses the same 75/25 "
     "correctness formula, so both arms produce comparable columns.", False),
])
add_para([
    ("5. Framework status. ", True),
    ("Every 0-1 value goes through framework_status in evaluation_policy.py: each metric "
     "has its own pass / review / fail bands (faithfulness passes at 0.70 and above, "
     "reviews down to 0.40), and tone and toxicity use a binary 0.5 cut-off.", False),
])
add_para([
    ("6. Reporting. ", True),
    ("Per-query rows are aggregated into session, sub-product, product and platform views "
     "- the decision levels the evaluation framework asks for.", False),
])

# ============================================================
# 12. Worked example
# ============================================================
add_heading("12. Worked example, step by step: JEV vs GPT on one unfaithful answer")
add_para([
    ("A real row from the held-out test split (session SYN-0001, turn 1), with the actual "
     "numbers saved in results/test_run. The bot's answer contains an invented benefit, so "
     "the designer label for faithfulness is ", False),
    ("fail", True),
    (".", False),
])
add_table(
    ["Field", "Text"],
    [
        ["User query", "How many casual leave days do employees receive each year?"],
        ["ADI response",
         "\u201cEmployees receive 12 days of Casual Leave per calendar year, pro-rated on "
         "joining. Unused Casual Leave can also be encashed at year end.\u201d"],
        ["Retrieved context",
         "The HR policy says 12 days per year, and that Casual Leave is neither carried "
         "forward nor encashed."],
        ["SME verdict",
         "No - the encashment claim is invented. Gold faithfulness label: fail."],
    ],
)
add_para([("Step by step with GPT-4o (RAGAS-style arm)", True)])
add_para([
    ("1. ", True),
    ("The prompt asks GPT to split the response into factual claims. GPT finds two: "
     "\u201c12 days per year\u201d and \u201cunused CL can be encashed at year end\u201d.", False),
])
add_para([
    ("2. ", True),
    ("GPT checks each claim against the retrieved context: claim 1 is supported; claim 2 "
     "is not (the policy says no encashment).", False),
])
add_para([
    ("3. ", True),
    ("GPT writes JSON: 1 claim supported out of 2. Our code computes faithfulness = "
     "1 \u00f7 2 = 0.50.", False),
])
add_para([
    ("4. ", True),
    ("Framework bands (pass \u2265 0.70, review \u2265 0.40): 0.50 \u2192 review. Under the "
     "older flat 0.5 cut-off this would have counted as a pass - one reason the "
     "metric-specific bands matter.", False),
])
add_para([("Step by step with JEV", True)])
add_para([
    ("1. ", True),
    ("Our code sends the row as state plus one Score question with four fixed levels. "
     "Level 1 says: \u201csome claims are supported, but at least one key fact is missing "
     "from or contradicts the context\u201d.", False),
])
add_para([
    ("2. ", True),
    ("JEV writes no text. It returns a probability for each level plus a confidence:", False),
])
add_table(
    ["Level", "Meaning (short)", "Probability"],
    [
        ["0", "Most claims missing from / contradict context", "0.01"],
        ["1", "Some supported, but a key fact contradicts the context", "0.99"],
        ["2", "Only a minor, non-essential detail is missing", "0.00"],
        ["3", "Every claim supported (or no factual claims)", "0.00"],
    ],
)
add_para([
    ("JEV is 99% sure this row sits on level 1 - exactly what was planted in the row.", False),
])
add_para([
    ("3. ", True),
    ("Code converts the answer to a number: raw = 0\u00d70.01 + 1\u00d70.99 + 2\u00d70 + "
     "3\u00d70 = 0.99, then score = 0.99 \u00f7 3 = 0.33.", False),
])
add_para([
    ("4. ", True),
    ("Framework bands: 0.33 < 0.40 \u2192 fail. This matches the designer label.", False),
])
add_para([("Side by side", True)])
add_table(
    ["Step", "RAGAS-style GPT-4o", "Matched-rubric GPT-4o", "JEV"],
    [
        ["How it sees the row", "1 of 2 claims supported", "Rubric level 1: a key fact "
         "contradicts the context", "Rubric level 1, probability 0.99"],
        ["Faithfulness score", "0.50", "0.33", "0.33"],
        ["Framework verdict", "review", "fail", "fail"],
        ["Confidence", "not reported", "not reported", "0.99"],
    ],
)
add_para([
    ("Why the numbers differ, in plain language. ", True),
    ("Both judges detected the same defect - the invented encashment benefit. They just "
     "express it differently. GPT must summarise \u201chow wrong\u201d the answer is as one "
     "ratio, and it gives the one correct claim credit, so the row lands at 0.50 = review. "
     "JEV only has to choose which pre-written level describes the row, and an invented "
     "benefit lands it on level 1 almost surely, so the score fails. The rubric, not the "
     "model, is what changes the number - that is exactly why the matched-rubric arm "
     "exists: when GPT-4o is handed JEV's levels, it also returns 0.33.", False),
])
add_para([
    ("Bonus: how JEV builds answer correctness on the same row. ", True),
    ("JEV answered three component questions: factual precision 0.6733, factual recall "
     "0.9967, semantic similarity 0.70. Code then computes F1 = 2 \u00d7 0.6733 \u00d7 "
     "0.9967 \u00f7 (0.6733 + 0.9967) = 0.8037, and answer correctness = 0.75 \u00d7 0.8037 "
     "+ 0.25 \u00d7 0.70 = 0.7778 \u2192 review. The gold label is 0.8268 (pass): JEV is "
     "slightly stricter because of the wrong extra sentence. On this row JEV lands closer "
     "to gold than the RAGAS arm (0.6667), while across the dataset GPT-4o's correctness "
     "grading is the closer of the two (MAE 0.041 vs 0.064), as described in section 7.2.", False),
])

# ============================================================
# 13. Beginner-friendly model explanations
# ============================================================
add_heading("13. How GPT and JEV work under the hood (beginner friendly)")
add_para([
    ("GPT-4o - an autocomplete engine that writes essays. ", True),
    ("GPT predicts the next word (token), over and over, based on everything written so "
     "far. It has no calculator and no memory between calls; everything it \u201cknows\u201d "
     "about the task comes from the prompt we send. As a judge, we give it instructions "
     "plus the row, and it types out a JSON answer. Our code parses that JSON and does all "
     "the arithmetic. Setting temperature = 0 makes it pick the most likely token each "
     "time, so repeat runs are stable. Analogy: an essay examiner who writes comments and "
     "a mark. Strengths: flexible, reads anything, produces human-readable reasons that "
     "auditors and SMEs can read. Weaknesses: the written reason and the number can "
     "disagree, the score is a summary rather than a distribution, it is slower and more "
     "expensive, and free text can occasionally fail to parse.", False),
])
add_para([
    ("JEV - a structured decision maker, not a writer. ", True),
    ("JEV (TypeSafe's System One model) never generates sentences. Every question has "
     "fixed, ordered answer options, and JEV returns a probability for each option plus "
     "one overall confidence. The score is a weighted average: raw = sum of level \u00d7 "
     "probability, and confidence measures how spread out the probabilities are "
     "({0.01, 0.99, 0, 0} = very sure; {0.24, 0.76, 0, 0} = less sure). Because the answer "
     "schema is fixed, JEV cannot hallucinate the output format. Analogy: a multiple-choice "
     "marker that ticks boxes and reports how sure it is about each tick. Strengths: fast, "
     "cheap, repeatable, and calibrated probabilities we can threshold on. Weaknesses: no "
     "written explanations, and everything depends on how well we write the answer levels - "
     "rubric design becomes the real work.", False),
])
add_para([
    ("One-line summary. ", True),
    ("GPT = flexible essay grader (text in, text out, we compute the score). JEV = fixed "
     "multiple-choice scorer (typed questions in, probabilities out, we normalise). This "
     "study's answer: JEV for high-volume routine metric scoring, GPT where written "
     "reasons matter.", False),
])

# ============================================================
# Move everything before Appendix A
# ============================================================
for el in created:
    body.remove(el)
for i, el in enumerate(created):
    body.insert(anchor_idx + i, el)

doc.save(DOCX)
print("saved ->", DOCX, "| elements inserted:", len(created))
